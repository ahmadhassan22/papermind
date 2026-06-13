from pathlib import Path
from langchain_community.document_loaders import PyPDFLoader
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_chroma import Chroma
from langchain_core.embeddings import Embeddings
import requests
from app.core.config import CHROMA_DIR, JINA_API_KEY

class JinaEmbeddings(Embeddings):
    def __init__(self):
        self.api_key = JINA_API_KEY
        self.url = "https://api.jina.ai/v1/embeddings"
        self.model = "jina-embeddings-v3"
        self.headers = {
            "Authorization": f"Bearer {self.api_key}",
            "Content-Type": "application/json"
        }

    def embed_documents(self, texts: list[str]) -> list[list[float]]:
        response = requests.post(
            self.url,
            headers=self.headers,
            json={"input": texts, "model": self.model}
        )
        response.raise_for_status()
        data = response.json()
        return [item["embedding"] for item in data["data"]]

    def embed_query(self, text: str) -> list[float]:
        return self.embed_documents([text])[0]

_embedding_model = None

def get_embedding_model():
    global _embedding_model
    if _embedding_model is None:
        _embedding_model = JinaEmbeddings()
    return _embedding_model

def load_document(file_path: Path) -> list:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        loader = PyPDFLoader(str(file_path))
        return loader.load()
    elif suffix == ".docx":
        import docx
        doc = docx.Document(str(file_path))
        text = "\n".join([para.text for para in doc.paragraphs if para.text.strip()])
        return [Document(page_content=text, metadata={"source": str(file_path), "page": 0})]
    elif suffix in [".md", ".txt"]:
        text = file_path.read_text(encoding="utf-8-sig")
        return [Document(page_content=text, metadata={"source": str(file_path), "page": 0})]
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

def ingest_document(file_path: Path) -> None:
    docs = load_document(file_path)
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200
    )
    chunks = splitter.split_documents(docs)
    for chunk in chunks:
        chunk.metadata["filename"] = Path(file_path).name

    embedding = get_embedding_model()
    Chroma.from_documents(
        documents=chunks,
        embedding=embedding,
        persist_directory=str(CHROMA_DIR)
    )